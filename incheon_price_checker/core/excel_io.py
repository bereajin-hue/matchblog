"""
인천e몰 파트너시스템 엑셀(21컬럼, MyData 시트) 읽기 / 결과 쓰기.
"""
from __future__ import annotations

from typing import Any

import re

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from config import COLOR_APPROVE, COLOR_HOLD, COLOR_REJECT, VERDICT_APPROVE, VERDICT_HOLD, VERDICT_REJECT

# 21컬럼 인덱스 (0-based)
_COL = {
    "상세보기": 0, "판매수": 1, "메인이미지": 2, "옵션": 3,
    "상품코드": 4, "상품명": 5, "시중가격": 6, "공급가": 7,
    "공급배송비": 8, "총공급가": 9, "판매가": 10, "배송비": 11,
    "총판매가": 12, "마진율": 13, "최저가": 14, "할인율": 15,
    "상태": 16, "진열": 17, "품절": 18, "판매자명": 19, "담당MD명": 20,
}

_RESULT_COLS = [
    "AI최저가", "AI단위최저가", "AI판정", "AI신뢰도", "AI근거",
    "재계산_할인율", "재계산_마진율",
    "동일후보1", "동일후보2", "동일후보3",
]

_FILL = {
    VERDICT_APPROVE: PatternFill("solid", fgColor=COLOR_APPROVE),
    VERDICT_HOLD:    PatternFill("solid", fgColor=COLOR_HOLD),
    VERDICT_REJECT:  PatternFill("solid", fgColor=COLOR_REJECT),
}


_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _safe_sheet_title(name: str) -> str:
    """엑셀 시트명 불가 문자 제거 후 31자 제한."""
    cleaned = _INVALID_SHEET_CHARS.sub("", name).strip() or "시트"
    return cleaned[:31]


def _to_int(value: Any) -> int:
    try:
        return int(value) if value is not None else 0
    except (ValueError, TypeError):
        return 0


def load_products(filepath: str) -> list[dict[str, Any]]:
    """
    MyData 시트에서 21컬럼을 읽어 상품 dict 리스트로 반환.
    헤더 행(1행)은 건너뛰고 2행부터 읽는다.
    """
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    if "MyData" not in wb.sheetnames:
        raise ValueError(f"'MyData' 시트를 찾을 수 없습니다. 시트 목록: {wb.sheetnames}")
    ws = wb["MyData"]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # 헤더로 컬럼 인덱스 동적 확인 (명세와 다를 경우 대비)
    header = list(rows[0])
    col_idx: dict[str, int] = {}
    for expected_col in _COL:
        try:
            col_idx[expected_col] = header.index(expected_col)
        except ValueError:
            col_idx[expected_col] = _COL[expected_col]  # fallback: 명세 순서

    products = []
    for row in rows[1:]:
        if not any(row):  # 빈 행 스킵
            continue
        def g(col: str) -> Any:
            idx = col_idx.get(col, _COL.get(col, -1))
            return row[idx] if 0 <= idx < len(row) else None

        products.append({
            "상세보기":   g("상세보기"),
            "판매수":    _to_int(g("판매수")),
            "메인이미지": g("메인이미지"),
            "옵션":      g("옵션"),
            "상품코드":   str(g("상품코드")) if g("상품코드") is not None else "",
            "상품명":    str(g("상품명")) if g("상품명") is not None else "",
            "시중가격":   _to_int(g("시중가격")),
            "공급가":    _to_int(g("공급가")),
            "공급배송비": _to_int(g("공급배송비")),
            "총공급가":   _to_int(g("총공급가")),
            "판매가":    _to_int(g("판매가")),
            "배송비":    _to_int(g("배송비")),
            "총판매가":   _to_int(g("총판매가")),
            "마진율":    _to_int(g("마진율")),
            "최저가":    _to_int(g("최저가")),
            "할인율":    _to_int(g("할인율")),
            "상태":      g("상태"),
            "진열":      g("진열"),
            "품절":      g("품절"),
            "판매자명":   g("판매자명"),
            "담당MD명":   g("담당MD명"),
        })

    wb.close()
    return products


def save_results(
    products: list[dict[str, Any]],
    results: list[dict[str, Any]],
    output_path: str,
    split_by_md: bool = False,
) -> None:
    """
    원본 21컬럼 보존 + 검증 컬럼(AI최저가 등 10개)을 우측에 추가해 저장.
    판정별 셀 색상 적용. split_by_md=True이면 담당MD명별로 시트 분리.
    """
    # 상품코드 → 결과 매핑
    result_map: dict[str, dict[str, Any]] = {
        str(r.get("상품코드", "")): r for r in results
    }

    headers_orig = list(_COL.keys())
    all_headers = headers_orig + _RESULT_COLS

    def make_sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
        # 헤더
        ws.append(all_headers)
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        for prod in rows:
            orig_vals = [prod.get(h) for h in headers_orig]
            res = result_map.get(str(prod.get("상품코드", "")), {})
            verdict = res.get("verdict", "")

            # 동일후보 최대 3개를 "상품명(가격)" 형태로
            top3 = res.get("top_candidates", [])
            cand_vals = []
            for i in range(3):
                if i < len(top3):
                    c = top3[i]
                    cand_vals.append(f"{c.get('title','')} ({c.get('lprice','')}원)")
                else:
                    cand_vals.append("")

            result_vals = [
                res.get("ai_lowest_price", ""),
                res.get("ai_unit_lowest_price", ""),
                verdict,
                res.get("confidence", ""),
                res.get("reason", ""),
                res.get("recalc_discount_rate", ""),
                res.get("recalc_margin_rate", ""),
            ] + cand_vals

            row_vals = orig_vals + result_vals
            ws.append(row_vals)

            # 판정 컬럼(AI판정, 3번째 결과 컬럼)에 색상 적용
            if verdict and verdict in _FILL:
                fill = _FILL[verdict]
                verdict_col = len(headers_orig) + 3  # AI판정 위치 (1-based)
                ws.cell(ws.max_row, verdict_col).fill = fill

        # 컬럼 너비 자동 조정 (상품명 컬럼 고정 넓게)
        ws.column_dimensions[get_column_letter(headers_orig.index("상품명") + 1)].width = 40
        ws.column_dimensions[get_column_letter(len(headers_orig) + 5)].width = 35  # AI근거

    wb_out = openpyxl.Workbook()

    if split_by_md:
        # 담당MD명별 시트 분리
        from collections import defaultdict
        md_groups: dict[str, list] = defaultdict(list)
        for prod in products:
            md = str(prod.get("담당MD명") or "미지정").strip()
            md_groups[md].append(prod)

        first = True
        for md_name, md_prods in md_groups.items():
            safe_title = _safe_sheet_title(md_name)
            if first:
                ws = wb_out.active
                ws.title = safe_title
                first = False
            else:
                ws = wb_out.create_sheet(title=safe_title)
            make_sheet(ws, md_prods)
    else:
        ws = wb_out.active
        ws.title = "결과"
        make_sheet(ws, products)

    wb_out.save(output_path)
